"""Check local questionnaire exports for omitted core and supplementary text."""
from pathlib import Path
import json,re,unicodedata
from docx import Document
root=Path(__file__).resolve().parents[1]
forms=json.loads((root/'audit/forms-snapshot.json').read_text())
def norm(s):return ''.join(c for c in unicodedata.normalize('NFKC',s)if not c.isspace())
results=[]
for f in forms:
 p=root/'audit/exports'/(f['id']+'-questionnaire.docx')
 if not p.exists():continue
 doc=Document(p);text=norm('\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells]))
 missing=[i['number']for i in f['items']if norm(i['text'])not in text]
 extraMissing=[i['text']for i in f.get('supplementaryItems',[])if norm(i['text'])not in text or any(norm(o)not in text for o in i['options'])]
 results.append({'id':f['id'],'missingItems':missing,'missingSupplementary':extraMissing})
(root/'audit/export-text-checks.json').write_text(json.dumps(results,ensure_ascii=False,indent=2))
if any(r['missingItems']or r['missingSupplementary']for r in results):raise SystemExit('Missing export content; inspect report')
print(f'Checked {len(results)} local questionnaire exports; no omitted core or supplementary text. Visual review remains separate.')
